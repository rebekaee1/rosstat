from app.data.cpi_provenance import cpi_provenance
from app.api.export import ExportIn, ExportMeta, _resolve_meta, _meta_rows


def test_reconstructed_cpi_export_preserves_origin_and_precision():
    notice = cpi_provenance('cpi-yoy')
    body = ExportIn(format='csv', filename='cpi.csv', points=[{'date': '2026-07-01', 'actual': 6.0}], meta=ExportMeta(
        indicator_name='ИПЦ год к году', unit='%', source='Росстат', provenance=notice))
    meta = _resolve_meta(body)
    rows = dict(_meta_rows(meta, 'Значение'))
    assert rows['Источник'] == 'Росстат'
    assert rows['Единица'] == '%'
    assert rows['Расчёт'] == notice
    assert 'округления' in notice
    assert cpi_provenance('cpi') is None
    assert cpi_provenance('key-rate') is None
    assert 'Calculated' in cpi_provenance('cpi-food-yoy', 'en')


def test_unknown_export_source_is_not_falsely_eurostat():
    rows = dict(_meta_rows({}, 'Value'))
    assert rows['Источник'] == 'Не указан'


def test_cpi_provenance_survives_csv_and_xlsx_http_export(auth_client, monkeypatch):
    from io import BytesIO
    from openpyxl import load_workbook
    from app.config import settings
    monkeypatch.setattr(settings, 'download_anon_limit', 2)
    notice = cpi_provenance('cpi-yoy')
    body = dict(filename='cpi', points=[dict(date='2026-07-01', actual=6.0)],
                meta=dict(indicator_name='ИПЦ год к году', unit='%', source='Росстат', provenance=notice))
    csv = auth_client.post('/api/v1/export/table', json={**body, 'format': 'csv'})
    assert csv.status_code == 200
    assert notice in csv.content.decode('utf-8-sig')
    xlsx = auth_client.post('/api/v1/export/table', json={**body, 'format': 'xlsx'})
    assert xlsx.status_code == 200
    wb = load_workbook(BytesIO(xlsx.content), read_only=True)
    rows = dict(wb['Описание'].iter_rows(values_only=True))
    assert rows['Расчёт'] == notice
    assert rows['Единица'] == '%'
