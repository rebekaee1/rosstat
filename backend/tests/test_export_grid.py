"""POST /export/grid: общая рамка выгрузки сетки (сравнение, рейтинг, регион, платежи)."""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.config import settings

URL = "/api/v1/export/grid"

BODY = {
    "format": "csv",
    "filename": "compare.csv",
    "title": "Россия и Турция",
    "columns": [
        {"key": "date", "label": "Дата"},
        {"key": "ru", "label": "Россия", "unit": "%"},
        {"key": "tr", "label": "Турция", "unit": "%"},
    ],
    "rows": [
        {"date": "2019-01-01", "ru": 4.5, "tr": 19.67},
        {"date": "2024-01-01", "ru": 7.4, "tr": None},
        {"date": "2025-01-01", "ru": 1234.5, "tr": 2},
    ],
    "meta": {"source": "Росстат; TÜİK", "country": "Россия, Турция", "frequency": "год"},
}


def register(tc, email="g@example.com"):
    r = tc.post("/api/v1/auth/register", json={"email": email, "password": "supersecret1", "consent": True})
    assert r.status_code == 201, r.text


def lines(resp):
    return resp.content.decode("utf-8-sig").split("\n")


def test_guest_blocked_by_default_limit(auth_client):
    r = auth_client.post(URL, json=BODY)
    assert r.status_code == 403 and r.json()["detail"]["code"] == "download_limit"


def test_guest_quota_shared_with_table(auth_client, monkeypatch):
    monkeypatch.setattr(settings, "download_anon_limit", 2)
    monkeypatch.setattr(settings, "download_anon_history_years", 0)
    r1 = auth_client.post(URL, json=BODY)
    assert r1.status_code == 200 and r1.headers["x-download-remaining"] == "1"
    table = {"format": "csv", "filename": "t.csv", "points": [{"date": "2020-01-01", "actual": 1}]}
    assert auth_client.post("/api/v1/export/table", json=table).status_code == 200   # общий счётчик
    r3 = auth_client.post(URL, json=BODY)
    assert r3.status_code == 403 and r3.json()["detail"]["code"] == "download_limit"


def test_csv_for_signed_in_user(auth_client):
    register(auth_client)
    r = auth_client.post(URL, json=BODY)
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert "attachment" in r.headers["content-disposition"]
    ls = lines(r)
    assert "# Таблица;Россия и Турция" in ls and any(l.startswith("# Источник;") for l in ls)
    assert any(l.startswith("# Страна;") for l in ls)
    head = ls.index("Дата;Россия, %;Турция, %")
    assert ls[head + 1] == "2019-01-01;4,5;19,67"
    assert ls[head + 2] == "2024-01-01;7,4;"          # пустая ячейка — пусто, не 0
    assert ls[head + 3].startswith("2025-01-01;1") and ";2" in ls[head + 3]


def test_xlsx_numbers_are_native(auth_client):
    register(auth_client)
    r = auth_client.post(URL, json={**BODY, "format": "xlsx", "filename": "compare.xlsx"})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    wb = load_workbook(BytesIO(r.content), read_only=True)
    assert wb.sheetnames == ["Описание", "Данные"]
    rows = list(wb["Данные"].iter_rows(values_only=True))
    assert rows[0] == ("Дата", "Россия, %", "Турция, %")
    assert rows[1] == ("2019-01-01", 4.5, 19.67) and rows[2][2] is None
    meta = " ".join(f"{a} {b}" for a, b in wb["Описание"].iter_rows(values_only=True) if a)
    assert "Росстат" in meta


def test_english_locale_labels(auth_client):
    register(auth_client)
    r = auth_client.post(URL, json=BODY, headers={"X-FE-Locale": "en"})
    ls = lines(r)
    assert "# Table;Россия и Турция" in ls and any(l.startswith("# Source;") for l in ls)
    assert "2019-01-01;4.5;19.67" in ls


@pytest.mark.parametrize("patch", [
    {"columns": [{"key": "lower", "label": "Нижняя"}, {"key": "date", "label": "Дата"}]},
    {"columns": [{"key": "UPPER", "label": "Верхняя"}]},
    {"rows": [{"date": "2020", "lower": 1.0}]},
    {"rows": [{"date": "2020", "Upper": 1.0}]},
    {"rows": []},
    {"columns": []},
    {"columns": [{"key": "a", "label": "A"}, {"key": "a", "label": "B"}]},
    {"columns": [{"key": "bad key", "label": "A"}]},
    {"columns": [{"key": "a", "label": ""}]},
    {"format": "pdf"},
    {"filename": ""},
])
def test_validation_and_no_ranges(auth_client, patch):
    register(auth_client)
    assert auth_client.post(URL, json={**BODY, **patch}).status_code == 422


def test_limits_on_size(auth_client, monkeypatch):
    register(auth_client)
    from app.services import export_grid
    monkeypatch.setattr(export_grid, "MAX_ROWS", 2)
    assert auth_client.post(URL, json=BODY).status_code == 422
    monkeypatch.setattr(export_grid, "MAX_ROWS", 50_000)
    monkeypatch.setattr(export_grid, "MAX_CELLS", 5)
    assert auth_client.post(URL, json=BODY).status_code == 422


def test_formula_injection_is_neutralised(auth_client):
    register(auth_client)
    body = {**BODY, "rows": [{"date": "=HYPERLINK(\"http://evil\",\"x\")", "ru": "-5,2", "tr": "+cmd|' /C calc'!A0"}],
            "title": "@SUM(1)"}
    csv_text = auth_client.post(URL, json=body).content.decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text and "'+cmd" in csv_text and "# Таблица;'@SUM(1)" in csv_text
    assert ";-5,2;" in csv_text                           # обычное отрицательное число не ломаем
    xl = auth_client.post(URL, json={**body, "format": "xlsx", "filename": "x.xlsx"})
    rows = list(load_workbook(BytesIO(xl.content), read_only=True)["Данные"].iter_rows(values_only=True))
    assert rows[1][0].startswith("'=") and rows[1][2].startswith("'+")


def test_guest_history_depth_follows_download_gate(auth_client, monkeypatch):
    monkeypatch.setattr(settings, "download_anon_limit", 5)
    monkeypatch.setattr(settings, "download_anon_history_years", 3)
    text = auth_client.post(URL, json=BODY).content.decode("utf-8-sig")
    assert "2019-01-01" not in text and "2024-01-01" in text and "2025-01-01" in text
    register(auth_client, "full@example.com")
    assert "2019-01-01" in auth_client.post(URL, json=BODY).content.decode("utf-8-sig")


def test_history_written_for_signed_in_when_cabinet_on(auth_client, monkeypatch):
    monkeypatch.setattr(settings, "cabinet_enabled", True)
    register(auth_client)
    body = {**BODY, "history": {"source": "compare", "subject_key": "ru-vs-tr", "params": {"codes": ["a", "b"], "rep": "yoy"}}}
    assert auth_client.post(URL, json=body).status_code == 200
    e = auth_client.get("/api/v1/cabinet/exports").json()["items"][0]
    assert e["source"] == "compare" and e["subject_key"] == "ru-vs-tr" and e["rows_count"] == 3
    assert e["params"]["codes"] == ["a", "b"] and e["params"]["filename"] == "compare.csv"
    assert e["params"]["columns"] == ["Дата", "Россия", "Турция"]
    assert "rows" not in e["params"]


def test_existing_table_endpoint_unchanged_for_old_clients(auth_client):
    register(auth_client)
    r = auth_client.post("/api/v1/export/table", json={"format": "csv", "filename": "a.csv", "points": [{"date": "2020-01-01", "actual": 1}]})
    assert r.status_code == 200
